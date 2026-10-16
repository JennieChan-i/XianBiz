import hashlib
import io
import json
import sqlite3
import zipfile
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image
from backend import files, service as s
from backend.app import create_app
from backend.db import BusinessError, Database, row
from conftest import order

def xlsx(kind,rows):
    wb=load_workbook(io.BytesIO(files.template(kind)));ws=wb.worksheets[0]
    for r in rows: ws.append(r)
    return files.excel_bytes(wb)

def test_import_preview_then_atomic_apply(db):
    content=xlsx('products',[['X1','服装','黑','M','',2500,10,1000,5,'上衣','品牌','秋'],['X1','服装','白','M','12345',2500,4,1200,2,'上衣','品牌','秋']])
    with db.transaction() as c:
        preview=files.preview_import(c,'products',content)
        assert not preview['errors'] and preview['total']==2
        assert c.execute('SELECT COUNT(*) FROM skus').fetchone()[0]==0
        result=files.apply_import(c,preview['id']);assert result['count']==2
        assert c.execute('SELECT COUNT(*) FROM products').fetchone()[0]==1
        assert row(c,'skus',1)['stock']==10 and row(c,'skus',2)['value']==4800
        with pytest.raises(BusinessError): files.apply_import(c,preview['id'])

def test_import_duplicate_and_negative_fields_rejected(db):
    content=xlsx('customer',[['A','100','','',200],['A','200','','',100],['B','','','',-1]])
    with db.transaction() as c:
        r=files.preview_import(c,'customer',content)
        assert r['total']==1 and len(r['errors'])==2
        with pytest.raises(BusinessError): files.apply_import(c,r['id'])
        assert c.execute('SELECT COUNT(*) FROM parties').fetchone()[0]==0

def test_import_formula_rejected(db):
    content=xlsx('customer',[['=1+1','','','',0]])
    with db.transaction() as c:
        r=files.preview_import(c,'customer',content)
        assert '公式' in r['errors'][0]['error']

def test_export_complete_and_excel_formula_escaped(setup):
    with setup['db'].transaction() as c:
        s.party(c,{'kind':'customer','name':'=HYPERLINK("http://evil")'})
        order(c,direct=True)
        wb=load_workbook(io.BytesIO(files.export_all(c)),data_only=False)
        assert set(files.EXPORTS)==set(wb.sheetnames)
        assert wb['客户供应商'].cell(5,3).value.startswith("'=")
        assert wb['销售采购'].max_row==2

def test_backup_restore_database_images_and_settings(setup):
    db=setup['db']; stream=io.BytesIO();Image.new('RGB',(200,100),'blue').save(stream,'PNG')
    im=files.save_image(db,1,stream.getvalue()); blob=files.backup(db)
    with db.transaction() as c:
        order(c,qty=2,direct=True);s.save_settings(c,{'shop_name':'变更后'})
    result=files.restore(db,blob,True)
    assert result['restored'] and __import__('pathlib').Path(result['protection']).exists()
    with db.connect() as c:
        assert row(c,'skus',1)['stock']==10 and s.settings(c)['shop_name']=='隔离测试'
        assert c.execute('SELECT COUNT(*) FROM orders').fetchone()[0]==0
    assert (db.directory/'images'/im['filename']).exists()

def test_backup_missing_image_cannot_silently_succeed(setup):
    db=setup['db']
    with db.transaction() as c:
        c.execute("INSERT INTO images(product_id,filename,thumbnail) VALUES(1,'missing.jpg','missing_t.jpg')")
    with pytest.raises(BusinessError,match='图片缺失'): files.backup(db)

def test_corrupt_backup_does_not_touch_current_database(setup):
    db=setup['db']; good=files.backup(db)
    out=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(good)) as a,zipfile.ZipFile(out,'w') as b:
        for name in a.namelist(): b.writestr(name,b'broken' if name=='database.sqlite3' else a.read(name))
    with pytest.raises(BusinessError): files.restore(db,out.getvalue(),True)
    with db.connect() as c: assert row(c,'skus',1)['stock']==10

def test_restore_requires_confirmation_and_blocks_zip_slip(setup):
    db=setup['db']; good=files.backup(db)
    with pytest.raises(BusinessError,match='确认'): files.restore(db,good,False)
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w') as z: z.writestr('../evil','x')
    with pytest.raises(BusinessError,match='不安全路径'): files.restore(db,out.getvalue(),True)

def test_future_schema_version_is_refused(db):
    with db.transaction() as c: c.execute("INSERT INTO schema_version VALUES(999,'future')")
    with pytest.raises(BusinessError,match='较新版本'): Database(db.directory)

@pytest.fixture
def client(tmp_path):
    app=create_app(tmp_path/'api');client=TestClient(app,base_url='http://127.0.0.1:9000')
    return client

def headers(client,key='test-command-key'):
    token=client.get('/api/session').json()['token']
    return {'X-XianBiz-Token':token,'X-Operation-Id':key}

def test_api_empty_bootstrap_no_login(client):
    r=client.get('/api/session');assert r.status_code==200 and r.json()['settings']=={}
    assert client.get('/api/dashboard').json()['today']['revenue']==0
    assert client.get('/api/health').json()['app']=='XianBiz'

@pytest.mark.parametrize('extra',[{'Origin':'https://evil.invalid'},{'Host':'evil.invalid'},{'Sec-Fetch-Site':'cross-site'},{}])
def test_write_security_origin_host_fetch_token(client,extra):
    r=client.post('/api/settings',json={'shop_name':'恶意'},headers=extra)
    assert r.status_code==403

def test_api_commands_idempotent_and_dates_readonly(client):
    p={'code':'A','name':'衣服','price':200,'skus':[{'color':'白','size':'均码','stock':10,'cost':100}]}
    h=headers(client)
    assert client.post('/api/products',json=p,headers=h).status_code==200
    assert client.post('/api/products',json=p,headers=h).status_code==200
    assert client.get('/api/skus').json()['total']==1
    assert client.patch('/api/orders/1',json={'created_at':'2000-01-01'},headers=h).status_code in (404,405)

def test_api_money_requires_integer_and_transaction_rollback(client):
    p={'kind':'customer','name':'A','opening':1.25}
    r=client.post('/api/parties',json=p,headers=headers(client));assert r.status_code==400
    assert client.get('/api/parties').json()['total']==0

def test_api_excel_and_backup_download(client):
    assert client.get('/api/export').content[:2]==b'PK'
    r=client.get('/api/backup');assert r.status_code==200
    assert client.post('/api/restore',files={'file':('backup.zip',r.content)},data={'confirmed':'true'},headers=headers(client)).status_code==200

def test_origin_blocked_even_with_token(client):
    h=headers(client);h['Origin']='http://127.0.0.1:9999'
    r=client.put('/api/settings',json={'shop_name':'A'},headers=h);assert r.status_code==403

def test_get_host_rebinding_rejected(client):
    assert client.get('/api/session',headers={'Host':'attacker.invalid'}).status_code==403
