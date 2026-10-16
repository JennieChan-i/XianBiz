import hashlib
import io
import json
import shutil
import sqlite3
import tempfile
import uuid
import zipfile
from contextlib import closing
from pathlib import Path, PurePosixPath
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageOps
from .db import VERSION, BusinessError, encode, insert, now, require
from . import service as s

HEADERS={
 'products': [('code','款号'),('name','商品名称'),('color','颜色'),('size','尺码'),('barcode','条码'),('price','批发价(分)'),('stock','期初件数'),('cost','期初成本(分)'),('threshold','库存阈值'),('category','分类'),('brand','品牌'),('season','季节')],
 'customer':[('name','名称'),('phone','电话'),('address','地址'),('note','备注'),('opening','期初应收(分)')],
 'supplier':[('name','名称'),('phone','电话'),('address','地址'),('note','备注'),('opening','期初应付(分)')]
}

def style_sheet(ws):
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]: cell.font=Font(color='FFFFFF',bold=True); cell.fill=PatternFill('solid',fgColor='234969')
    for column in ws.columns:
        width=min(45,max(14,max((len(str(c.value or '')) for c in list(column)[:200]),default=10)+3))
        ws.column_dimensions[column[0].column_letter].width=width

def safe_excel(value):
    if isinstance(value,str) and value.startswith(('=','+','-','@')): return "'"+value
    return value

def excel_bytes(wb):
    stream=io.BytesIO(); wb.save(stream); return stream.getvalue()

def template(kind):
    require(kind in HEADERS,'不支持的模板类型')
    wb=Workbook(); ws=wb.active; ws.title='期初导入'
    ws.append([label for key,label in HEADERS[kind]]); style_sheet(ws)
    help=wb.create_sheet('填写说明')
    for line in ['仅导入期初资料，不导入历史交易。','金额单位为整数分，数量为整数件。','商品同款可多行，不同颜色尺码组合不可重复；条码可空。','期初库存和成本只允许对新SKU填写。','名称、款号、颜色、尺码为文本；请先预览校验再确认。']:
        help.append([line])
    help.column_dimensions['A'].width=90
    return excel_bytes(wb)

def preview_import(conn,kind,content):
    require(kind in HEADERS,'导入类型无效'); require(len(content)<=15*1024*1024,'Excel文件不能超过15MB')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            require(sum(i.file_size for i in z.infolist())<=100*1024*1024,'Excel解压内容过大')
        wb=load_workbook(io.BytesIO(content),read_only=True,data_only=False)
    except BusinessError: raise
    except Exception: raise BusinessError('无法读取Excel，请使用下载的xlsx模板')
    sheet=wb.worksheets[0]; it=sheet.iter_rows(values_only=True)
    columns=next(it,None); require(columns is not None and list(columns)==[label for k,label in HEADERS[kind]],'模板列不匹配，请重新下载模板')
    rows=[]; errors=[]; keys=set(); bars=set()
    numeric={'price','stock','cost','threshold','opening'}
    for idx,values in enumerate(it,2):
        require(idx<=10001,'单次导入最多10000行')
        if not any(v is not None for v in values): continue
        p={}
        try:
            for (key,label),value in zip(HEADERS[kind],values):
                if isinstance(value,str): require(not value.startswith('='),'不接受公式，请粘贴为数值')
                if key in numeric:
                    if key=='stock': require(value is not None and value!='','期初件数为必填，无库存请填写0')
                    v=value if value is not None else (5 if key=='threshold' else 0)
                    if isinstance(v,float) and v.is_integer(): v=int(v)
                    p[key]=s.integer(v,label)
                else: p[key]=s.text(str(value) if value is not None else '',label,key in ('code','name','color','size'))
            key=(p['code'],p['color'],p['size']) if kind=='products' else p['name']
            require(key not in keys,'文件内款色码或名称重复'); keys.add(key)
            if kind=='products':
                require(not conn.execute('SELECT 1 FROM skus s JOIN products p ON p.id=s.product_id WHERE p.code=? AND s.color=? AND s.size=?',key).fetchone(),'该SKU已存在，不能重复导入期初')
                if p['barcode']:
                    require(p['barcode'] not in bars and not conn.execute('SELECT 1 FROM skus WHERE barcode=?',(p['barcode'],)).fetchone(),'条码重复'); bars.add(p['barcode'])
            else:
                require(not conn.execute('SELECT 1 FROM parties WHERE kind=? AND name=?',(kind,p['name'])).fetchone(),'往来名称已存在')
            rows.append(p)
        except BusinessError as error: errors.append({'row':idx,'error':error.message})
    wb.close()
    require(bool(rows) or bool(errors),'Excel没有可导入行')
    id=uuid.uuid4().hex
    insert(conn,'import_jobs',id=id,kind=kind,payload=encode(rows),errors=encode(errors),created_at=now())
    return {'id':id,'rows':rows[:200],'total':len(rows),'errors':errors,'preview_limit':200}

def apply_import(conn,id):
    job=conn.execute('SELECT * FROM import_jobs WHERE id=?',(id,)).fetchone()
    require(job and not job['applied'],'预览不存在或已导入'); require(not json.loads(job['errors']),'存在错误，不能导入')
    rows=json.loads(job['payload'])
    if job['kind']=='products':
        groups={}
        for r in rows: groups.setdefault(r['code'],[]).append(r)
        for code,items in groups.items():
            old=conn.execute('SELECT * FROM products WHERE code=?',(code,)).fetchone()
            p=dict(old) if old else {key:items[0][key] for key in ('code','name','category','brand','season','price')}
            for item in items:
                require(not conn.execute('SELECT 1 FROM skus s JOIN products p ON p.id=s.product_id WHERE p.code=? AND s.color=? AND s.size=?',(code,item['color'],item['size'])).fetchone(),'预览后SKU发生变化，请重新预览')
            p['skus']=items; p['inventory_mode']='opening'; s.product(conn,p,old['id'] if old else None)
    else:
        for r in rows: s.party(conn,{**r,'kind':job['kind']})
    conn.execute('UPDATE import_jobs SET applied=1 WHERE id=?',(id,))
    return {'count':len(rows)}

EXPORTS={'商品':'products','SKU库存':'skus','商品图片索引':'images','客户供应商':'parties','销售采购':'orders','订单明细':'order_lines','履约':'fulfillments','退换货':'returns','收付款':'funds','统一实际资金流水':'cash_ledger','渠道明细':'fund_channels','债项':'debts','贷项预收付':'credits','核销':'allocations','退款分配':'refund_allocations','经营费用':'expenses','库存流水':'inventory_ledger','成本调整':'cost_adjustments','负库存待核成本':'negative_lots','盘亏待核成本':'inventory_shortages','预留变动':'reservation_events','损益流水':'profit_ledger','盘点单':'stocktakes','实盘记录':'stocktake_items','操作记录':'history','店铺配置':'settings'}

def export_all(conn):
    wb=Workbook(); wb.remove(wb.active)
    for title,table in EXPORTS.items():
        ws=wb.create_sheet(title); cur=conn.execute(f'SELECT * FROM {table}')
        ws.append([d[0] for d in cur.description])
        for r in cur: ws.append([safe_excel(v) for v in r])
        style_sheet(ws)
    return excel_bytes(wb)

def export_report(conn,start,end):
    from .queries import report
    r=report(conn,start,end); wb=Workbook(); ws=wb.active; ws.title='经营汇总'
    ws.append(['项目','金额(分)']);
    for key in ('revenue','original_cost','cost_adjustment','gross_profit','inventory_profit','expenses','profit','purchase'): ws.append([key,r[key]])
    style_sheet(ws)
    for key,title in [('details','损益明细'),('ranking','销量排行'),('trend','日期统计')]:
        ws=wb.create_sheet(title)
        if r[key]:
            ws.append(list(r[key][0]));
            for line in r[key]: ws.append([safe_excel(v) for v in line.values()])
            style_sheet(ws)
    return excel_bytes(wb)

def save_image(db,product_id,content):
    require(len(content)<=15*1024*1024,'图片不能超过15MB')
    try:
        im=Image.open(io.BytesIO(content)); require(im.width*im.height<=30_000_000,'图片尺寸过大')
        im=ImageOps.exif_transpose(im).convert('RGB')
        im.thumbnail((1600,1600))
    except BusinessError: raise
    except Exception: raise BusinessError('图片格式不支持或已损坏')
    with db.lock:
        filename=uuid.uuid4().hex+'.jpg'; thumb='thumb_'+filename
        directory=db.directory/'images'
        try:
            im.save(directory/filename,quality=85,optimize=True); im.thumbnail((240,240)); im.save(directory/thumb,quality=80)
            with db.transaction() as conn:
                s.row(conn,'products',product_id)
                pos=conn.execute('SELECT COUNT(*) FROM images WHERE product_id=?',(product_id,)).fetchone()[0]
                require(pos<12,'每款最多12张图片')
                id=insert(conn,'images',product_id=product_id,filename=filename,thumbnail=thumb,position=pos)
            return {'id':id,'filename':filename}
        except BaseException:
            (directory/filename).unlink(missing_ok=True); (directory/thumb).unlink(missing_ok=True); raise

def backup(db):
    with db.lock,tempfile.TemporaryDirectory() as temp:
        snapshot=Path(temp)/'database.sqlite3'
        with db.connect() as conn,closing(sqlite3.connect(snapshot)) as dest: conn.backup(dest)
        data={'database.sqlite3':snapshot.read_bytes()}
        with db.connect() as conn:
            for r in conn.execute('SELECT filename,thumbnail FROM images'):
                for name in r:
                    path=db.directory/'images'/name
                    require(path.is_file(),'商品图片缺失，备份已停止：'+name,409)
                    data['images/'+name]=path.read_bytes()
            data['settings.json']=encode(s.settings(conn)).encode('utf-8')
        manifest={'format':'XianBiz.full','version':VERSION,'created_at':now(),'files':{name:hashlib.sha256(b).hexdigest() for name,b in data.items()}}
        out=io.BytesIO()
        with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
            for name,b in data.items(): z.writestr(name,b)
            z.writestr('manifest.json',encode(manifest))
        return out.getvalue()

def restore(db,content,confirmed):
    require(confirmed is True,'恢复会替换当前经营数据，必须明确确认')
    require(len(content)<=500*1024*1024,'备份超过500MB，请联系维护人员离线迁移')
    with db.lock,tempfile.TemporaryDirectory() as temp:
        base=Path(temp)
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                names=z.namelist()
                require(len(names)==len(set(names)) and len(names)<=50000,'备份条目重复或过多')
                require(sum(i.file_size for i in z.infolist())<=2*1024**3,'备份展开内容过大')
                for name in names:
                    path=PurePosixPath(name)
                    require(not path.is_absolute() and '..' not in path.parts and '\\' not in name and ':' not in name,'备份包含不安全路径')
                manifest=json.loads(z.read('manifest.json'))
                require(manifest['format']=='XianBiz.full' and manifest['version']==VERSION,'备份格式或版本不兼容')
                require(set(names)==set(manifest['files'])|{'manifest.json'},'备份清单不匹配')
                require('database.sqlite3' in names and 'settings.json' in names,'缺少数据库或配置')
                for name,digest in manifest['files'].items():
                    data=z.read(name); require(hashlib.sha256(data).hexdigest()==digest,'备份校验失败：'+name)
                    path=base/name; path.parent.mkdir(parents=True,exist_ok=True); path.write_bytes(data)
            staged=base/'database.sqlite3'
            with closing(sqlite3.connect(staged)) as conn:
                require(conn.execute('PRAGMA integrity_check').fetchone()[0]=='ok' and not conn.execute('PRAGMA foreign_key_check').fetchall(),'备份数据库损坏')
                require(conn.execute('SELECT MAX(version) FROM schema_version').fetchone()[0]==VERSION,'数据库版本不兼容')
                for r in conn.execute('SELECT filename,thumbnail FROM images'):
                    for name in r: require('images/'+name in manifest['files'],'备份商品图片缺失')
            protection=db.directory/'recovery'; protection.mkdir(exist_ok=True)
            protect=protection/('恢复前_'+now().replace(':','-')+'.zip')
            protect.write_bytes(backup(db))
            target=db.directory/'images'; old=base/'old_images'
            # All validations happen before touching live data. Writers share this lock.
            target.rename(old)
            try:
                shutil.copytree(base/'images',target) if (base/'images').exists() else target.mkdir()
                with closing(sqlite3.connect(staged)) as source,db.connect() as dest: source.backup(dest)
            except BaseException:
                if target.exists(): shutil.rmtree(target)
                old.rename(target)
                raise
            return {'restored':True,'protection':str(protect)}
        except BusinessError: raise
        except Exception as error: raise BusinessError('恢复失败，备份格式无效或文件损坏') from error
